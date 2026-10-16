"""Build testdata/device: a fake laptop home with every kind of file that Raven meets.

It starts from the same home as make_testdata.py (documents, photos, a video with speech), then adds:
office files (PPTX, XLSX, ODT, EPUB, HTML, RTF), a scanned PDF, photos with camera EXIF, screenshots,
WebP, GIF, HEIC, JPEG XL, AVIF, SVG, TIFF, BMP, near-duplicate and exact-duplicate photos, videos in
MP4, WebM, MKV and MOV, audio in MP3, M4A, WAV, OGG, FLAC and Opus (speech and a tone), archives (zip, 7z,
tar.gz, tar.zst), a code project with .git and node_modules, an SDK, a dataset, installers, a disk image,
a SQLite file, a font, secrets, a hidden .ssh folder, a partial download, a large file, a symbolic link,
an empty folder, deep folders and names with spaces and Unicode.

Each file that Raven must find gets a unique phrase. device_manifest.json lists the checks that
run_device_test.py runs after the index is complete.

Usage: python testdata/make_device.py [out_dir] [--real-videos]   (default testdata/device; it needs ffmpeg, espeak-ng,
cjxl, heif-enc, magick, 7z, zstd, soffice and tesseract, and it skips a part whose tool is missing)
"""

from __future__ import annotations

import json
import os
import random
import shutil
import subprocess
import sys
import tarfile
import tempfile
import zipfile
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, str(Path(__file__).parent))
import make_testdata as base  # noqa: E402

HERE = Path(__file__).parent
CHECKS: list[dict] = []
SKIPPED: list[str] = []


def check(kind: str, **kw) -> None:
    CHECKS.append({"kind": kind, **kw})


def have(tool: str) -> bool:
    if shutil.which(tool):
        return True
    SKIPPED.append(tool)
    return False


def run(*cmd: str) -> None:
    subprocess.run(cmd, check=True, capture_output=True)


def font(size: int):
    for f in ("DejaVuSans.ttf", "/usr/share/fonts/TTF/DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(f, size)
        except OSError:
            pass
    return ImageFont.load_default()


def text_image(path: Path, lines: list[str], size=(900, 600), bg="white", fg="black") -> Path:
    im = Image.new("RGB", size, bg)
    d = ImageDraw.Draw(im)
    for i, ln in enumerate(lines):
        d.text((40, 40 + i * 56), ln, fill=fg, font=font(34))
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path)
    return path


def scene(path: Path, sky, sun, ground, size=(1600, 1200), shift=0, exif_camera: str | None = None, quality=92) -> Path:
    im = Image.new("RGB", size, sky)
    d = ImageDraw.Draw(im)
    w, h = size
    d.ellipse([w // 3 + shift, h // 5, w // 3 + w // 4 + shift, h // 5 + w // 4], fill=sun)
    d.rectangle([0, h * 2 // 3, w, h], fill=ground)
    for i in range(12):
        d.line([(i * w // 12, h * 2 // 3), (i * w // 12 + 40, h)], fill=(40, 40, 40), width=3)
    path.parent.mkdir(parents=True, exist_ok=True)
    kw = {"quality": quality} if path.suffix.lower() in (".jpg", ".jpeg") else {}
    if exif_camera:
        ex = im.getexif()
        ex[271], ex[272] = exif_camera.split(" ", 1)
        kw["exif"] = ex
    im.save(path, **kw)
    return path


def speech(path: Path, text: str) -> Path | None:
    if not (have("espeak-ng") and have("ffmpeg")):
        return None
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as t:
        wav = Path(t) / "s.wav"
        run("espeak-ng", "-s", "150", "-w", str(wav), text)
        if path.suffix == ".wav":
            shutil.copy(wav, path)
        else:
            run("ffmpeg", "-nostdin", "-v", "error", "-y", "-i", str(wav), str(path))
    return path


# ---------------------------------------------------------------- parts


def documents(home: Path) -> None:
    d = home / "Documents" / "Office"
    d.mkdir(parents=True, exist_ok=True)
    from openpyxl import Workbook
    from pptx import Presentation

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Zephyrine solar kiosk pitch"
    s.placeholders[1].text = "The zephyrine kiosk charges forty phones a day from one panel."
    p.save(d / "kiosk_pitch.pptx")
    check("find", query="zephyrine solar kiosk", expect="kiosk_pitch.pptx", what="PPTX text")

    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "Quokka lantern budget"])
    ws.append(["Lanterns", 1200])
    ws.append(["Batteries", 340])
    wb.save(d / "club_budget.xlsx")
    check("find", query="quokka lantern budget", expect="club_budget.xlsx", what="XLSX cells")

    base.make_docx(d / "marigold_minutes.docx", "# Marigold committee minutes\n\nThe marigold committee chose Saturday for the cleanup drive.\n")
    check("find", query="marigold committee cleanup drive", expect="marigold_minutes.docx", what="DOCX text")
    if have("soffice"):
        with tempfile.TemporaryDirectory() as t:
            run("soffice", f"-env:UserInstallation=file://{t}/p", "--headless", "--convert-to", "odt", "--outdir", str(d),
                str(d / "marigold_minutes.docx"))
        (d / "marigold_minutes.odt").rename(d / "heron_notes.odt")
        check("find", query="marigold committee", expect="heron_notes.odt", what="ODT text")

    (d / "orchid_page.html").write_text("<html><head><title>Orchid repair guide</title></head><body><h1>Orchid repair guide</h1>"
                                         "<p>Tighten the orchid bracket with a 4 mm hex key.</p></body></html>")
    check("find", query="orchid bracket hex key", expect="orchid_page.html", what="HTML text")
    (d / "notes.rtf").write_text(r"{\rtf1\ansi Pelican ferry timetable: the pelican ferry leaves at 7 am.}")
    check("name", query="notes.rtf", expect="notes.rtf", what="RTF by name")

    # EPUB: a zip with a mimetype file, a container and one XHTML chapter.
    ep = d / "lighthouse_story.epub"
    with zipfile.ZipFile(ep, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                   '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        z.writestr("OEBPS/content.opf", '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
                   '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">x</dc:identifier><dc:title>The Lighthouse Keeper</dc:title>'
                   '<dc:language>en</dc:language></metadata><manifest><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/></manifest>'
                   '<spine><itemref idref="c1"/></spine></package>')
        z.writestr("OEBPS/c1.xhtml", '<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><body><h1>Chapter one</h1>'
                   "<p>The lighthouse keeper fed the tangerine cat every evening.</p></body></html>")
    check("find", query="lighthouse keeper tangerine cat", expect="lighthouse_story.epub", what="EPUB text")

    # A scanned PDF: the page is a picture of text, with no text layer. OCR must read it.
    img = text_image(Path(tempfile.mkdtemp()) / "scan.png", ["RENTAL RECEIPT", "Received from Kiran Patil", "Saffron apartments rent paid",
                                                            "Amount Rs 14,500 for May"], size=(1240, 900))
    doc = pymupdf.open()
    page = doc.new_page(width=620, height=450)
    page.insert_image(page.rect, filename=str(img))
    doc.save(home / "Documents" / "scan_rent_receipt.pdf")
    check("find", query="saffron apartments rent", expect="scan_rent_receipt.pdf", what="scanned PDF (OCR)")

    (home / "Documents" / "plain notes with spaces.txt").write_text("Walrus pottery class starts on the third Monday.\n")
    check("find", query="walrus pottery class", expect="plain notes with spaces.txt", what="name with spaces")
    (home / "Documents" / "résumé_ünïcode.md").write_text("# Résumé\n\nNarwhal robotics intern, summer project on line following.\n")
    check("find", query="narwhal robotics intern", expect="résumé_ünïcode.md", what="Unicode name")
    deep = home / "Documents" / "a" / "b" / "c" / "d" / "e" / "f"
    deep.mkdir(parents=True, exist_ok=True)
    (deep / "deep_note.txt").write_text("The capybara spreadsheet is in the deepest folder.\n")
    check("find", query="capybara spreadsheet", expect="deep_note.txt", what="deep folder")
    (home / "Documents" / "Empty folder").mkdir(exist_ok=True)
    (home / "Documents" / "data.json").write_text(json.dumps({"project": "Albatross tracker", "owner": "team 7"}))
    check("find", query="albatross tracker", expect="data.json", what="JSON text")
    (home / "Documents" / "server.log").write_text("".join(f"2026-03-0{i % 9 + 1} INFO request ok\n" for i in range(200)) + "ERROR flamingo cache timeout\n")
    check("find", query="flamingo cache timeout", expect="server.log", what="log file")


def pictures(home: Path) -> None:
    p = home / "Pictures"
    cam = p / "Camera"
    scene(cam / "IMG_20250412_101500.jpg", (90, 150, 230), (255, 220, 60), (30, 140, 50), exif_camera="Canon EOS 200D")
    check("visual", query="a sunny landscape with a yellow sun", expect="IMG_20250412_101500.jpg", what="camera photo by appearance")
    check("no_ocr", expect="IMG_20250412_101500.jpg", what="camera photo skips OCR")
    # Near duplicates: the same scene, smaller and slightly moved.
    scene(cam / "IMG_20250412_101501.jpg", (90, 150, 230), (255, 220, 60), (30, 140, 50), size=(800, 600), exif_camera="Canon EOS 200D")
    scene(cam / "IMG_20250412_101502.jpg", (90, 150, 230), (255, 220, 60), (30, 140, 50), shift=8, exif_camera="Canon EOS 200D")
    check("near", expect=["IMG_20250412_101500.jpg", "IMG_20250412_101501.jpg", "IMG_20250412_101502.jpg"], what="similar photos")
    # An exact duplicate in another folder.
    shutil.copy(cam / "IMG_20250412_101500.jpg", p / "copy of trip photo.jpg")
    check("dupe", expect=["IMG_20250412_101500.jpg", "copy of trip photo.jpg"], what="exact duplicate")

    shots = p / "Screenshots"
    text_image(shots / "Screenshot_2025-06-01_upi.png", ["Payment successful", "Paid to Mango Grove Stores", "UPI ref 8812 3301 7720"])
    check("find", query="mango grove stores payment", expect="Screenshot_2025-06-01_upi.png", what="screenshot OCR")

    fmt = p / "Formats"
    fmt.mkdir(parents=True, exist_ok=True)
    scene(fmt / "night_sky.webp", (10, 10, 40), (240, 240, 255), (5, 5, 20))
    check("visual", query="a dark night sky with a moon", expect="night_sky.webp", what="WebP by appearance")
    frames = [Image.new("RGB", (200, 200), c) for c in ((255, 0, 0), (0, 255, 0), (0, 0, 255))]
    frames[0].save(fmt / "blink.gif", save_all=True, append_images=frames[1:], duration=200, loop=0)
    check("name", query="blink.gif", expect="blink.gif", what="GIF by name")
    scene(fmt / "forest.tiff", (150, 200, 150), (255, 255, 200), (20, 90, 30))
    scene(fmt / "desert.bmp", (250, 200, 120), (255, 240, 180), (210, 160, 80), size=(800, 600))
    check("visual", query="a photo of a sandy desert", expect="desert.bmp", what="BMP by appearance")
    (fmt / "logo.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" width="64" height="64"><rect width="64" height="64" fill="#0b7a55"/></svg>')
    check("name", query="logo.svg", expect="logo.svg", what="SVG by name")
    src = scene(Path(tempfile.mkdtemp()) / "red_car.png", (200, 200, 210), (200, 30, 30), (80, 80, 80))
    if have("cjxl"):
        run("cjxl", str(src), str(fmt / "red_car.jxl"), "--quiet")
        check("visual", query="a grey road scene with a red circle", expect="red_car.jxl", what="JPEG XL by appearance")
    if have("heif-enc"):
        run("heif-enc", str(src), "-o", str(fmt / "phone_photo.heic"))
        check("name", query="phone_photo.heic", expect="phone_photo.heic", what="HEIC by name")
    if have("magick"):
        run("magick", str(src), str(fmt / "web_banner.avif"))
        check("name", query="web_banner.avif", expect="web_banner.avif", what="AVIF by name")
    # A chart: search by what it shows.
    im = Image.new("RGB", (900, 600), "white")
    d = ImageDraw.Draw(im)
    for i, hgt in enumerate([120, 300, 220, 420, 360]):
        d.rectangle([100 + i * 140, 520 - hgt, 180 + i * 140, 520], fill=(40, 110, 200))
    d.line([(80, 520), (820, 520)], fill="black", width=3)
    im.save(fmt / "quarterly.png")
    check("visual", query="a blue bar chart", expect="quarterly.png", what="chart by appearance")


def media(home: Path) -> None:
    v = home / "Videos"
    a = home / "Music"
    rec = home / "Recordings"
    if have("ffmpeg"):
        v.mkdir(parents=True, exist_ok=True)
        tmp = Path(tempfile.mkdtemp())
        green = scene(tmp / "g.png", (120, 200, 120), (250, 250, 250), (20, 120, 20), size=(640, 480))
        night = scene(tmp / "n.png", (10, 10, 40), (240, 240, 255), (5, 5, 20), size=(640, 480))
        talk = speech(tmp / "talk.wav", "Welcome to the kayak safety briefing. Always wear the orange life jacket on the river.")
        args = ["ffmpeg", "-nostdin", "-v", "error", "-y", "-loop", "1", "-t", "5", "-i", str(green), "-loop", "1", "-t", "5", "-i", str(night)]
        chain = "[0:v]fps=25,setsar=1[a];[1:v]fps=25,setsar=1[b];[a][b]concat=n=2:v=1:a=0[out]"
        aud = ["-i", str(talk), "-map", "2:a", "-c:a", "aac"] if talk else []
        subprocess.run([*args, *aud[:2], "-filter_complex", chain, "-map", "[out]", *aud[2:], "-c:v", "libx264", "-g", "50",
                        "-pix_fmt", "yuv420p", str(v / "kayak_briefing.mp4")], check=True, capture_output=True)
        if talk:
            check("find", query="kayak safety orange life jacket", expect="kayak_briefing.mp4", what="speech in MP4 (transcript)")
        check("visual", query="a dark night sky", expect="kayak_briefing.mp4", what="video moment", at_min=4.0)
        for ext, codec in (("webm", ["-c:v", "libvpx-vp9", "-b:v", "200k"]), ("mkv", ["-c:v", "libx264"]), ("mov", ["-c:v", "libx264"])):
            subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-y", "-loop", "1", "-t", "3", "-i", str(green), *codec, "-pix_fmt", "yuv420p",
                            str(v / f"clip.{ext}")], check=True, capture_output=True)
            check("name", query=f"clip.{ext}", expect=f"clip.{ext}", what=f"{ext.upper()} video by name")
        # Music: a tone, no speech. The transcript must stay empty.
        a.mkdir(parents=True, exist_ok=True)
        run("ffmpeg", "-nostdin", "-v", "error", "-y", "-f", "lavfi", "-i", "sine=frequency=440:duration=8", str(a / "tone_440.flac"))
        check("no_speech", expect="tone_440.flac", what="music has no transcript")
    for ext in ("mp3", "m4a", "wav", "ogg", "opus"):
        speech(rec / f"memo_{ext}.{ext}", f"Voice memo in {ext}. Buy the cardamom and the teal notebook before Sunday.")
    if (rec / "memo_mp3.mp3").exists():
        # "sunday": the robot voice of espeak-ng makes Whisper mishear rare words such as "cardamom".
        check("transcripts", expect=[f"memo_{e}.{e}" for e in ("mp3", "m4a", "wav", "ogg", "opus")], word="sunday",
              what="speech in MP3, M4A, WAV, OGG, Opus")


def archives(home: Path) -> None:
    d = home / "Downloads" / "Archives"
    d.mkdir(parents=True, exist_ok=True)
    src = Path(tempfile.mkdtemp()) / "trip_docs"
    (src / "inner").mkdir(parents=True)
    (src / "itinerary.txt").write_text("Gondola ride booked for the second day.\n")
    (src / "inner" / "tickets.md").write_text("# Tickets\nGondola tickets in the inner folder.\n")
    with zipfile.ZipFile(d / "trip_docs.zip", "w") as z:
        for f in src.rglob("*"):
            z.write(f, f.relative_to(src.parent))
    with tarfile.open(d / "trip_docs.tar.gz", "w:gz") as t:
        t.add(src, "trip_docs")
    if have("zstd"):
        subprocess.run(f"tar -C {src.parent} -cf - trip_docs | zstd -q -o {d / 'trip_docs.tar.zst'}", shell=True, check=True)
    if have("7z"):
        run("7z", "a", "-y", str(d / "trip_docs.7z"), str(src))
    check("archive", expect=["trip_docs.zip", "trip_docs.tar.gz", "trip_docs.tar.zst", "trip_docs.7z"], member="trip_docs/inner/tickets.md",
          what="archives open like folders")


def software(home: Path) -> None:
    # A code project: names only. node_modules and .git are never read.
    proj = home / "Codes" / "weather-app"
    (proj / "src").mkdir(parents=True, exist_ok=True)
    (proj / ".git").mkdir(exist_ok=True)
    (proj / ".git" / "HEAD").write_text("ref: refs/heads/main\n")
    (proj / "package.json").write_text('{"name": "weather-app", "version": "1.0.0"}')
    (proj / "src" / "forecast.js").write_text("// The pangolin forecast module.\nexport const pangolin = 1\n")
    nm = proj / "node_modules" / "left-pad"
    nm.mkdir(parents=True, exist_ok=True)
    (nm / "index.js").write_text("module.exports = 1\n")
    check("names_only", expect="forecast.js", word="pangolin", what="code project: name only")
    check("absent", expect="node_modules/left-pad/index.js", what="node_modules not listed")
    # An SDK and a dataset: the folder only, not its contents.
    sdk = home / "Android" / "Sdk" / "platforms" / "android-34"
    sdk.mkdir(parents=True, exist_ok=True)
    (home / "Android" / "Sdk" / "source.properties").write_text("Pkg.Revision=34\n")
    for i in range(300):
        (sdk / f"res_{i}.xml").write_text(f"<r id='{i}'/>")
    check("absent", expect="Android/Sdk/platforms/android-34/res_7.xml", what="SDK contents not listed")
    ds = home / "Codes" / "learning" / "scikit_learn_data" / "faces"
    ds.mkdir(parents=True, exist_ok=True)
    for i in range(200):
        Image.new("RGB", (64, 64), (i, 100, 200 - i % 200)).save(ds / f"face_{i}.jpg")
    check("absent", expect="scikit_learn_data/faces/face_5.jpg", what="dataset contents not listed")
    # Installers, a disk image, a database, a font.
    dl = home / "Downloads"
    (dl / "spotify_1.2.deb").write_bytes(b"!<arch>\n" + os.urandom(4096))
    (dl / "Obsidian-1.5.AppImage").write_bytes(b"\x7fELF" + os.urandom(8192))
    (dl / "setup.exe").write_bytes(b"MZ" + os.urandom(4096))
    (dl / "ubuntu-mini.iso").write_bytes(os.urandom(64 * 1024))
    check("category", expect="spotify_1.2.deb", category="Installer", what="installer category")
    import sqlite3

    con = sqlite3.connect(home / "Documents" / "contacts.db")
    con.execute("create table c(name)")
    con.commit()
    con.close()
    check("name", query="contacts.db", expect="contacts.db", what="SQLite file by name")
    fonts = [f for f in ("/usr/share/fonts/TTF/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf") if os.path.exists(f)]
    if fonts:
        (home / ".fonts").mkdir(exist_ok=True)
        shutil.copy(fonts[0], home / "Documents" / "DejaVuSans.ttf")
        check("category", expect="DejaVuSans.ttf", category="Font", what="font category")


def private(home: Path) -> None:
    ssh = home / ".ssh"
    ssh.mkdir(exist_ok=True)
    (ssh / "id_ed25519").write_text("-----BEGIN OPENSSH PRIVATE KEY-----\nnot a real key\n-----END OPENSSH PRIVATE KEY-----\n")
    check("absent", expect=".ssh/id_ed25519", what=".ssh never indexed")
    (home / "Documents" / "passwords.txt").write_text("bank: hunter2\nwifi: Saffron@2024\n")
    check("category", expect="passwords.txt", category="Secret", what="secret category")
    (home / "Documents" / "kyc.txt").write_text("Name: Kiran Patil\nPAN: ABCDE1234F\nAadhaar: 2345 6789 0123\nPhone: 9876543210\n")
    check("pii", expect="kyc.txt", what="PII flagged")


def space(home: Path) -> None:
    dl = home / "Downloads"
    with open(dl / "big_backup.bin", "wb") as f:
        f.write(os.urandom(48 * 1024 * 1024))
    check("name", query="big_backup.bin", expect="big_backup.bin", what="large file by name")
    (dl / "movie.mp4.crdownload").write_bytes(os.urandom(2048))
    check("name", query="movie.mp4.crdownload", expect="movie.mp4.crdownload", what="partial download by name")
    os.symlink(dl / "big_backup.bin", dl / "link_to_backup")
    check("absent", expect="Downloads/link_to_backup", what="symbolic links skipped")


def unnamed(home: Path) -> None:
    """PLAN2 P2-3.2: files with generic names. "named" checks want words in the new name; "kind" checks a label."""
    u = home / "Downloads" / "Unnamed"
    u.mkdir(parents=True, exist_ok=True)
    from openpyxl import Workbook

    docs = [  # (name, title, body, words that the new name must hold)
        ("Untitled(1).pdf", "Hostel Fee Receipt", "Hostel fee of 42000 rupees received for the winter term.", ["hostel", "fee"]),
        ("download (7).pdf", "Internship Offer Letter", "We are glad to offer you the role of data intern.", ["internship", "offer"]),
        ("scan_0042.pdf", "Vehicle Insurance Policy", "Policy for the scooter, valid for one year.", ["insurance"]),
        ("file_example_PDF_1MB.pdf", "Library Membership Rules", "Members can borrow four books at a time.", ["library"]),
        ("1.pdf", "Electricity Bill March", "Units used: 212. Amount due: 1890.", ["electricity"]),
        ("a8f3c2e1b9.pdf", "Bus Pass Renewal Form", "Fill in the route and the student number.", ["bus", "pass"]),
        ("3f2a9c1e-7b4d-4e8a-9c2b-1a2b3c4d5e6f.pdf", "Lab Safety Checklist", "Wear goggles and gloves at all times.", ["lab", "safety"]),
    ]
    for name, title, body, words in docs:
        base.make_pdf(u / name, [f"{title}\n\n{body}"])
        check("named", expect=name, words=words, what=f"name for {name}")
    base.make_docx(u / "Doc1.docx", "# Robotics Club Agenda\n\nLine follower demo and budget review on Friday.\n")
    check("named", expect="Doc1.docx", words=["robotics"], what="name for Doc1.docx")
    base.make_docx(u / "New Microsoft Word Document.docx", "# Trip Packing List\n\nTorch, raincoat, power bank.\n")
    check("named", expect="New Microsoft Word Document.docx", words=["packing"], what="name for a Word default name")
    wb = Workbook()
    wb.active["A1"] = "Monthly Grocery Budget"
    wb.active["A2"] = "Rice"
    wb.save(u / "Book1.xlsx")
    check("named", expect="Book1.xlsx", words=["grocery"], what="name for Book1.xlsx")
    (u / "New Text Document.txt").write_text("Wedding Gift Ideas\nA lamp, a book, a plant.\n")
    check("named", expect="New Text Document.txt", words=["wedding"], what="name for a text default name")
    (u / "output.txt").write_text("Server Migration Notes\nMove the database on Sunday night.\n")
    check("named", expect="output.txt", words=["migration"], what="name for output.txt")
    (u / "document(2).md").write_text("# Thesis Chapter Outline\n\nIntroduction, method, results.\n")
    check("named", expect="document(2).md", words=["thesis"], what="name for a Markdown file")
    shots = [
        ("Screenshot_2025-07-14-09-12-33.png", ["Order placed", "Blue Heron Bookstore", "Total 640"], ["heron", "2025-07-14"]),
        ("Screenshot from 2025-08-02 10-22-11.png", ["Boarding pass", "Indigo Flight 6E 211", "Gate 14"], ["boarding", "2025-08-02"]),
        ("Screenshot 2025-09-05 at 10.22.11 AM.png", ["Exam timetable", "Physics on Monday"], ["exam", "2025-09-05"]),
        ("IMG-20250610-WA0003.jpg", ["Rent paid", "Saffron Apartments", "June"], ["rent", "2025-06-10"]),
        ("WhatsApp Image 2025-05-01 at 18.20.11.jpeg", ["Library fine notice", "Return the books"], ["library", "2025-05-01"]),
        ("image.png", ["Wifi password card", "Network Kestrel Home"], ["wifi"]),
    ]
    for name, lines, words in shots:
        text_image(u / name, lines)
        check("named", expect=name, words=words, what=f"name for {name}")
    photos = [
        ("IMG_20250412_183000.jpg", ["2025-04-12"]), ("DSC_0001.JPG", []), ("PXL_20250301_101500123.jpg", ["2025-03-01"]),
        ("20250622_071500.jpg", ["2025-06-22"]), ("photo (4).jpg", []),
    ]
    for name, words in photos:
        scene(u / name, (250, 140, 60), (255, 230, 120), (40, 40, 80), size=(800, 600))
        check("named", expect=name, words=words, what=f"name for photo {name}")
    for name, text in (("Recording (3).m4a", "Reminder to call the plumber about the kitchen sink on Tuesday."),
                       ("voice_001.ogg", "Shopping list for the picnic: bread, cheese and grapes."),
                       ("audio (2).mp3", "Notes for the chemistry viva: titration and the burette.")):
        if speech(u / name, text):
            check("named", expect=name, words=[], what=f"name for audio {name}")
    if have("ffmpeg"):
        img = scene(Path(tempfile.mkdtemp()) / "v.png", (120, 200, 120), (250, 250, 250), (20, 120, 20), size=(320, 240))
        for name in ("VID_20250105_120000.mp4", "WhatsApp Video 2025-02-03 at 09.00.00.mp4"):
            run("ffmpeg", "-nostdin", "-v", "error", "-y", "-loop", "1", "-t", "2", "-i", str(img), "-c:v", "libx264",
                "-pix_fmt", "yuv420p", str(u / name))
        check("named", expect="VID_20250105_120000.mp4", words=["2025-01-05"], what="name for a camera video")
        check("named", expect="WhatsApp Video 2025-02-03 at 09.00.00.mp4", words=["2025-02-03"], what="name for a WhatsApp video")
    for name in ("Final Report.pdf", "CS3401 Notes.pdf"):
        base.make_pdf(u / name, [f"{name[:-4]}\n\nA file with a real name."])
        check("not_named", expect=name, what=f"keep the name {name}")

    # Kind checks (P2-3.1): the content decides when the name says nothing or the wrong thing.
    k = home / "Downloads" / "Kinds"
    k.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (64, 64), "teal").save(k / "picture", "PNG")
    check("kind", expect="Downloads/Kinds/picture", label="PNG", what="PNG with no extension")
    base.make_pdf(k / "notes.txt", ["A PDF with a text name"])
    check("kind", expect="Downloads/Kinds/notes.txt", label="Text", mismatch=True, what="PDF named .txt: mismatch")
    (k / "invoice.pdf").write_bytes(b"MZ\x90\x00" + os.urandom(2000))
    check("kind", expect="Downloads/Kinds/invoice.pdf", label="PDF", mismatch=True, warning=True, what="program named .pdf: warning")
    (k / "Makefile").write_text("all:\n\techo hi\n")
    check("kind", expect="Downloads/Kinds/Makefile", label="Makefile", what="Makefile")
    with tarfile.open(k / "backup.tar.gz", "w:gz") as t:
        t.add(k / "Makefile", arcname="Makefile")
    check("kind", expect="Downloads/Kinds/backup.tar.gz", label="Tar archive (gzip)", what=".tar.gz")


# Real videos for search by appearance and speech (--real-videos, needs yt-dlp and the network).
# Licenses: Big Buck Bunny is CC BY (Blender Foundation). The NASA videos are NASA media (generally public domain).
REAL_VIDEOS = [
    ("big_buck_bunny.mp4", "https://www.youtube.com/watch?v=aqz-KE-bpKQ"),
    ("perseverance_landing.mp4", "https://www.youtube.com/watch?v=rzmd7RouGrM"),
    ("nasa_hunt_for_asteroids.mp4", "https://www.youtube.com/watch?v=MQi0R6PCROI"),
]


def real_videos(home: Path) -> None:
    if not have("yt-dlp"):
        return
    v = home / "Videos"
    v.mkdir(parents=True, exist_ok=True)
    cache = HERE / ".video_cache"  # kept between runs, so a new device folder needs no new download
    cache.mkdir(exist_ok=True)
    for name, url in REAL_VIDEOS:
        if not (cache / name).exists():
            subprocess.run(["yt-dlp", "--no-warnings", "-q", "-f", "bv*[height<=480][ext=mp4]+ba[ext=m4a]/b[height<=480]",
                            "--merge-output-format", "mp4", "-o", str(cache / name), url], check=True)
        shutil.copy(cache / name, v / name)
    check("visual", query="a big white rabbit", expect="big_buck_bunny.mp4", what="real video: a rabbit")
    check("visual", query="a butterfly", expect="big_buck_bunny.mp4", what="real video: a butterfly")
    check("visual", query="a rocket landing on mars", expect="perseverance_landing.mp4", what="real video: a Mars landing")
    check("find", query="planetary defense near earth objects", expect="nasa_hunt_for_asteroids.mp4", what="real video: narration")
    check("ask", question="How is NASA planning to defend Earth against asteroids?", word="asteroid", what="Ask from real narration")
    check("no_speech", expect="perseverance_landing.mp4", what="music-only video: no transcript")


def asks(_: Path) -> None:
    check("ask", question="What is the late fee in my rent agreement?", word="500", what="Ask with a citation")
    check("ask", question="What did the kayak briefing say to wear?", word="jacket", what="Ask from a video transcript")
    check("ask_none", question="Who won the cricket world cup in 1983?", what="Ask says not found")


def main(argv: list[str]) -> int:
    args = [a for a in argv[1:] if not a.startswith("--")]
    out = (Path(args[0]) if args else HERE / "device").expanduser().resolve()
    random.seed(11)
    base.main(["make_testdata", str(out)])  # the documents and photos of the base test home
    base._scratch_root = None  # base.main removed its scratch folder: the next helper call makes a new one
    parts = [documents, pictures, media, archives, software, private, space, unnamed, asks]
    if "--real-videos" in argv:
        parts.append(real_videos)
    for part in parts:
        part(out)
    files = [p for p in out.rglob("*") if p.is_file() and not p.is_symlink()]
    (HERE / "device_manifest.json").write_text(json.dumps({"root": str(out), "checks": CHECKS, "skipped_tools": sorted(set(SKIPPED))}, indent=1,
                                                          ensure_ascii=False))
    print(f"device: {len(files)} files, {sum(p.stat().st_size for p in files) / 1e6:.0f} MB in {out}")
    print(f"checks: {len(CHECKS)} in {HERE / 'device_manifest.json'}")
    if SKIPPED:
        print(f"skipped (tool missing): {', '.join(sorted(set(SKIPPED)))}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
